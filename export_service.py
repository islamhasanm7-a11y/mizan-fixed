import html
"""
MizanPro - Executive Excel Financial Reporting (15 Sheets) & Unified Printable HTML Templates Service
النظام المالي المتكامل - ميزان برو
محرك تصدير حزمة التقارير المالية والزكوية والضريبية والمخزنية والأصول والرواتب وأعمار الديون (15 صفحة) + قوالب الطباعة
"""

import io
from datetime import date
from typing import Dict, Any, Optional
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from models import (
    JournalEntry, SalesInvoice, PurchaseInvoice, CreditDebitNote,
    CashVoucher, PayrollRun, CompanySettings, InvoiceType, NoteType, VoucherType
)
from ledger_service import LedgerService
from financial_engine import FinancialStatementsEngine
from zatca_phase2_validator import ZATCAPhase2Validator


class ExcelExportService:
    """محرك تصدير ملفات الإكسيل المالية والزكوية والضريبية والرقابية المعتمدة (15 صفحة)"""

    HEADER_FILL = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    SUBHEADER_FILL = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")
    SECTION_FILL = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    TOTAL_FILL = PatternFill(start_color="FEF08A", end_color="FEF08A", fill_type="solid")
    ACCENT_FILL = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")

    FONT_TITLE = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    FONT_HEADER = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    FONT_BOLD = Font(name="Calibri", size=11, bold=True, color="000000")
    FONT_REGULAR = Font(name="Calibri", size=10, color="000000")

    ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")
    ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
    ALIGN_CENTER = Alignment(horizontal="center", vertical="center")

    THIN_BORDER = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1")
    )
    DOUBLE_BOTTOM_BORDER = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="000000"),
        bottom=Side(style="double", color="000000")
    )

    @classmethod
    def generate_master_excel_report(
        cls,
        db: Session,
        company_name: Optional[str] = None,
        fiscal_year: int = 2026,
        start_date: Optional[date] = None,
        end_date: Optional[date] = None
    ) -> io.BytesIO:
        """توليد مصنف الإكسيل المالي والضريبي والمخزني والرقابي الشامل كاملاً من 15 صفحة"""
        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        settings = LedgerService.get_company_settings(db)
        company_name = company_name or settings.company_name_ar
        fiscal_year = fiscal_year or settings.fiscal_year

        tb_df = LedgerService.get_live_trial_balance(db, start_date=start_date, end_date=end_date, include_zero_balances=True)
        inv_df = LedgerService.get_inventory_valuation_df(db)
        fa_df = LedgerService.get_fixed_assets_df(db)
        emp_df = LedgerService.get_employees_and_eos_df(db)
        aging_df = LedgerService.get_aging_report(db, party_type="CUSTOMER")
        cc_df = LedgerService.get_cost_centers_pnl(db, start_date=start_date, end_date=end_date)

        engine = FinancialStatementsEngine(trial_balance_df=tb_df, fiscal_year=fiscal_year, vat_rate=settings.default_vat_rate)
        full_report = engine.generate_full_financial_report()

        cls._create_cover_sheet(wb, company_name, fiscal_year, full_report)
        cls._create_trial_balance_sheet(wb, tb_df)
        cls._create_income_statement_sheet(wb, company_name, fiscal_year, full_report["income_statement"])
        cls._create_balance_sheet_sheet(wb, company_name, fiscal_year, full_report["balance_sheet"])
        cls._create_cash_flow_sheet(wb, company_name, fiscal_year, full_report["cash_flow_statement"])
        cls._create_zakat_sheet(wb, company_name, fiscal_year, full_report["zakat_declaration"])
        cls._create_vat_return_sheet(wb, company_name, fiscal_year, full_report["vat_return"])
        cls._create_ratios_sheet(wb, company_name, fiscal_year, full_report["financial_ratios"])
        cls._create_dataframe_sheet(wb, "جرد وتقييم المخزون", inv_df)
        cls._create_dataframe_sheet(wb, "سجل الأصول الثابتة", fa_df)
        cls._create_dataframe_sheet(wb, "الموظفون ونهاية الخدمة", emp_df)
        cls._create_dataframe_sheet(wb, "أعمار ديون العملاء", aging_df)
        cls._create_dataframe_sheet(wb, "ربحية مراكز التكلفة", cc_df)
        cls._create_notes_sheet(wb, company_name, fiscal_year)
        cls._create_journal_log_sheet(wb, db, start_date=start_date, end_date=end_date)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @classmethod
    def _create_cover_sheet(cls, wb, company_name: str, year: int, report: Dict[str, Any]):
        ws = wb.create_sheet(title="الملخص التنفيذي")
        ws.views.sheetView[0].rightToLeft = True

        ws.merge_cells("B2:G2")
        cell = ws["B2"]
        cell.value = f"التقرير المالي والزكوي والضريبي والرقابي الشامل (15 صفحة) - {company_name}"
        cell.font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
        cell.fill = cls.HEADER_FILL
        cell.alignment = cls.ALIGN_CENTER

        ws["B4"] = "السنة المالية:"
        ws["C4"] = str(year)
        ws["B5"] = "المعايير المطبقة:"
        ws["C5"] = "المعايير الدولية للتقرير المالي (IFRS for SMEs) / الهيئة السعودية (SOCPA) / ZATCA"
        ws["B6"] = "تاريخ استخراج التقرير:"
        ws["C6"] = date.today().strftime("%Y-%m-%d")

        for r in range(4, 7):
            ws[f"B{r}"].font = cls.FONT_BOLD
            ws[f"C{r}"].font = cls.FONT_REGULAR

        pnl = report["income_statement"]
        bs = report["balance_sheet"]
        zakat = report["zakat_declaration"]
        vat = report["vat_return"]
        cf = report["cash_flow_statement"]

        kpi_data = [
            ("إجمالي الأصول (Total Assets)", bs["total_assets"], "0F172A"),
            ("صافي الإيرادات التشغيلية (Net Revenues)", pnl["net_revenue"], "1E3A8A"),
            ("صافي ربح الفترة (Net Profit)", pnl["net_profit"], "15803D"),
            ("رصيد النقدية والبنوك الختامي (Closing Cash)", cf["closing_cash_balance"], "0369A1"),
            ("صافي ضريبة القيمة المضافة المستحقة (Net VAT Due)", vat["net_vat_due"], "6D28D9"),
            ("وعاء الزكاة الشرعي (Zakat Base)", zakat["zakat_base"], "B45309"),
            ("الزكاة المستحقة (Zakat Due)", zakat["zakat_due"], "B91C1C"),
        ]

        row = 9
        for title, val, color in kpi_data:
            ws[f"B{row}"] = title
            ws[f"B{row}"].font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            ws[f"B{row}"].fill = PatternFill(start_color=color, fill_type="solid")
            ws[f"B{row}"].alignment = cls.ALIGN_RIGHT

            ws[f"C{row}"] = val
            ws[f"C{row}"].font = cls.FONT_BOLD
            ws[f"C{row}"].number_format = "#,##0.00"
            ws[f"C{row}"].alignment = cls.ALIGN_LEFT
            ws[f"C{row}"].border = cls.THIN_BORDER
            row += 2

        ws.column_dimensions["B"].width = 42
        ws.column_dimensions["C"].width = 30

    @classmethod
    def _create_trial_balance_sheet(cls, wb, tb_df: pd.DataFrame):
        ws = wb.create_sheet(title="ميزان المراجعة")
        ws.views.sheetView[0].rightToLeft = True

        headers = [
            "كود الحساب", "اسم الحساب (عربي)", "النوع", "التبويب الفرعي",
            "مدين أول المدة", "دائن أول المدة", "حركة مدين", "حركة دائن",
            "مدين آخر المدة", "دائن آخر المدة", "صافي الرصيد"
        ]

        for col_num, h_text in enumerate(headers, start=1):
            cell = ws.cell(row=2, column=col_num, value=h_text)
            cell.font = cls.FONT_HEADER
            cell.fill = cls.HEADER_FILL
            cell.alignment = cls.ALIGN_CENTER
            cell.border = cls.THIN_BORDER

        row_idx = 3
        for _, r in tb_df.iterrows():
            ws.cell(row=row_idx, column=1, value=r["Account Code"]).alignment = cls.ALIGN_CENTER
            ws.cell(row=row_idx, column=2, value=r["Account Name"]).alignment = cls.ALIGN_RIGHT
            ws.cell(row=row_idx, column=3, value=r["Account Type"]).alignment = cls.ALIGN_CENTER
            ws.cell(row=row_idx, column=4, value=r["Account SubType"]).alignment = cls.ALIGN_CENTER

            for c_idx, field in enumerate([
                "Opening Debit", "Opening Credit", "Period Debit", "Period Credit",
                "Closing Debit", "Closing Credit", "Net Balance"
            ], start=5):
                c_val = float(r.get(field, 0.0))
                c_cell = ws.cell(row=row_idx, column=c_idx, value=c_val)
                c_cell.number_format = "#,##0.00"
                c_cell.alignment = cls.ALIGN_LEFT

            for col in range(1, len(headers) + 1):
                ws.cell(row=row_idx, column=col).border = cls.THIN_BORDER
                ws.cell(row=row_idx, column=col).font = cls.FONT_REGULAR

            row_idx += 1

        ws.cell(row=row_idx, column=2, value="الإجمالي العام").font = cls.FONT_BOLD
        ws.cell(row=row_idx, column=2).alignment = cls.ALIGN_CENTER

        for c_idx, col_name in enumerate([
            "Opening Debit", "Opening Credit", "Period Debit", "Period Credit",
            "Closing Debit", "Closing Credit", "Net Balance"
        ], start=5):
            tot = float(tb_df[col_name].sum()) if not tb_df.empty else 0.0
            tot_cell = ws.cell(row=row_idx, column=c_idx, value=tot)
            tot_cell.font = cls.FONT_BOLD
            tot_cell.number_format = "#,##0.00"
            tot_cell.fill = cls.TOTAL_FILL
            tot_cell.border = cls.DOUBLE_BOTTOM_BORDER

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    @classmethod
    def _create_income_statement_sheet(cls, wb, company_name: str, year: int, pnl: Dict[str, Any]):
        ws = wb.create_sheet(title="قائمة الدخل الشامل")
        ws.views.sheetView[0].rightToLeft = True

        ws.merge_cells("B2:D2")
        ws["B2"] = f"قائمة الدخل الشامل للسنة المالية المنتهية في {year}/12/31"
        ws["B2"].font = cls.FONT_TITLE
        ws["B2"].fill = cls.HEADER_FILL
        ws["B2"].alignment = cls.ALIGN_CENTER

        lines = [
            ("إجمالي إيرادات المبيعات والخدمات", pnl["gross_sales"], False, False),
            ("يخصم: مردودات ومسموحات المبيعات", -pnl["sales_returns"], False, False),
            ("يخصم: خصم مسموح به (خصم مبيعات)", -pnl["sales_discounts"], False, False),
            ("صافي الإيرادات التشغيلية", pnl["net_revenue"], True, True),
            ("يخصم: صافي تكلفة البضاعة المباعة والتسويات المخزنية", -pnl["cost_of_sales"], False, False),
            ("مجمل الربح (Gross Profit)", pnl["gross_profit"], True, True),
            ("المصاريف التشغيلية والإدارية:", None, False, False),
            ("  - الرواتب والأجور وبدلات الموظفين", pnl["operating_expenses"]["salaries"], False, False),
            ("  - إيجار المقرات والمعارض والمستودعات", pnl["operating_expenses"]["rent"], False, False),
            ("  - المنافع (كهرباء، مياه، اتصالات)", pnl["operating_expenses"]["utilities"], False, False),
            ("  - الصيانة العامة وقطع الغيار", pnl["operating_expenses"]["maintenance"], False, False),
            ("  - الدعاية والتسويق", pnl["operating_expenses"]["marketing"], False, False),
            ("  - الرسوم الحكومية والتراخيص", pnl["operating_expenses"]["government_fees"], False, False),
            ("  - مصروف إهلاك الأصول الثابتة", pnl["operating_expenses"]["depreciation"], False, False),
            ("  - مصروف مكافأة نهاية الخدمة", pnl["operating_expenses"]["end_of_service"], False, False),
            ("  - مصاريف عمومية وإدارية أخرى", pnl["operating_expenses"]["other_admin"], False, False),
            ("إجمالي المصاريف التشغيلية والإدارية", -pnl["operating_expenses"]["total"], True, False),
            ("الربح التشغيلي (Operating Profit)", pnl["operating_profit"], True, True),
            ("يضاف: إيرادات وأرباح أخرى متنوعة", pnl["other_income"], False, False),
            ("يخصم: المصاريف والتكاليف التمويلية", -pnl["finance_costs"], False, False),
            ("صافي الربح قبل الزكاة الشرعية", pnl["profit_before_zakat"], True, True),
            ("يخصم: مخصص الزكاة الشرعية التقديري", -pnl["zakat_expense"], False, False),
            ("صافي ربح الفترة النهائي بعد الزكاة", pnl["net_profit"], True, True)
        ]

        curr_row = 4
        for title, val, is_bold, is_highlight in lines:
            ws.cell(row=curr_row, column=2, value=title).font = cls.FONT_BOLD if is_bold else cls.FONT_REGULAR
            ws.cell(row=curr_row, column=2).alignment = cls.ALIGN_RIGHT
            ws.cell(row=curr_row, column=2).border = cls.THIN_BORDER

            val_cell = ws.cell(row=curr_row, column=3, value=val if val is not None else "")
            val_cell.font = cls.FONT_BOLD if is_bold else cls.FONT_REGULAR
            val_cell.alignment = cls.ALIGN_LEFT
            val_cell.border = cls.DOUBLE_BOTTOM_BORDER if is_highlight else cls.THIN_BORDER

            if val is not None:
                val_cell.number_format = "#,##0.00;(#,##0.00);\"-\""

            if is_highlight:
                ws.cell(row=curr_row, column=2).fill = cls.ACCENT_FILL if "صافي ربح" in title or "مجمل" in title else cls.SECTION_FILL
                val_cell.fill = cls.ACCENT_FILL if "صافي ربح" in title or "مجمل" in title else cls.SECTION_FILL

            curr_row += 1

        ws.column_dimensions["B"].width = 48
        ws.column_dimensions["C"].width = 25

    @classmethod
    def _create_balance_sheet_sheet(cls, wb, company_name: str, year: int, bs: Dict[str, Any]):
        ws = wb.create_sheet(title="قائمة المركز المالي")
        ws.views.sheetView[0].rightToLeft = True

        ws.merge_cells("B2:D2")
        ws["B2"] = f"قائمة المركز المالي كما في {year}/12/31"
        ws["B2"].font = cls.FONT_TITLE
        ws["B2"].fill = cls.HEADER_FILL
        ws["B2"].alignment = cls.ALIGN_CENTER

        ca = bs["current_assets"]
        nca = bs["non_current_assets"]
        cl = bs["current_liabilities"]
        ncl = bs["non_current_liabilities"]
        eq = bs["equity"]

        lines = [
            ("الأصول (ASSETS)", None, True, "SECTION"),
            ("الأصول المتداولة (Current Assets):", None, True, None),
            ("  - النقد وما في حكمه بالبنوك والخزينة", ca["cash_and_equivalents"], False, None),
            ("  - العملاء والمدينون التجاريون (بالصافي)", ca["net_trade_receivables"], False, None),
            ("  - المخزون السلعي (بالمتوسط المرجح للتكلفة)", ca["inventory"], False, None),
            ("  - مصروفات مدفوعة مقدماً وأرصدة مدينة", ca["prepayments_and_other"], False, None),
            ("  - ضريبة القيمة المضافة - مدخلات مستردة", ca["input_vat_recoverable"], False, None),
            ("إجمالي الأصول المتداولة", ca["total"], True, "TOTAL"),

            ("الأصول غير المتداولة (Non-Current Assets):", None, True, None),
            ("  - الأصول الثابتة الملموسة (بالتكلفة)", nca["ppe_gross"], False, None),
            ("  - يخصم: مجمع إهلاك الأصول الثابتة", -nca["accumulated_depreciation"], False, None),
            ("  - صافي الأصول الثابتة الملموسة", nca["net_ppe"], True, None),
            ("  - الأصول غير الملموسة والبرمجيات", nca["intangibles"], False, None),
            ("إجمالي الأصول غير المتداولة", nca["total"], True, "TOTAL"),
            ("🔵 إجمالي الأصول (TOTAL ASSETS)", bs["total_assets"], True, "GRAND_TOTAL"),

            ("الالتزامات وحقوق الملكية (LIABILITIES & EQUITY)", None, True, "SECTION"),
            ("الالتزامات المتداولة (Current Liabilities):", None, True, None),
            ("  - الموردون والدائنون التجاريون", cl["trade_payables"], False, None),
            ("  - مصروفات مستحقة وأرصدة دائنة", cl["accruals_and_other"], False, None),
            ("  - ضريبة القيمة المضافة المستحقة للهيئة", cl["output_vat_payable"], False, None),
            ("  - مخصص الزكاة والضريبة المستحق", cl["zakat_provision_payable"], False, None),
            ("إجمالي الالتزامات المتداولة", cl["total"], True, "TOTAL"),

            ("الالتزامات غير المتداولة (طويلة الأجل):", None, True, None),
            ("  - قروض وتسهيلات تمويلية طويلة الأجل", ncl["long_term_loans"], False, None),
            ("  - مخصص مكافأة نهاية الخدمة للموظفين", ncl["eos_provision"], False, None),
            ("إجمالي الالتزامات غير المتداولة", ncl["total"], True, "TOTAL"),
            ("إجمالي الالتزامات (Total Liabilities)", bs["total_liabilities"], True, "TOTAL"),

            ("حقوق الملكية (EQUITY):", None, True, None),
            ("  - رأس المال المدفوع", eq["paid_up_capital"], False, None),
            ("  - الاحتياطي النظامي / القانوني", eq["statutory_reserve"], False, None),
            ("  - الأرباح المبقاة (المرحلة)", eq["retained_earnings_opening"], False, None),
            ("  - جاري الشركاء / المالك", eq["partners_current_account"], False, None),
            ("  - صافي ربح الفترة الحالي", eq["current_period_net_profit"], False, None),
            ("إجمالي حقوق الملكية (Total Equity)", eq["total"], True, "TOTAL"),
            ("🟢 إجمالي الالتزامات وحقوق الملكية", bs["total_liabilities_and_equity"], True, "GRAND_TOTAL")
        ]

        cls._write_styled_statement_lines(ws, lines)

    @classmethod
    def _create_cash_flow_sheet(cls, wb, company_name: str, year: int, cf: Dict[str, Any]):
        ws = wb.create_sheet(title="قائمة التدفقات النقدية")
        ws.views.sheetView[0].rightToLeft = True

        ws.merge_cells("B2:D2")
        ws["B2"] = f"قائمة التدفقات النقدية (الطريقة غير المباشرة IAS 7) لعام {year}"
        ws["B2"].font = cls.FONT_TITLE
        ws["B2"].fill = cls.HEADER_FILL
        ws["B2"].alignment = cls.ALIGN_CENTER

        op = cf["operating_activities"]
        inv = cf["investing_activities"]
        fin = cf["financing_activities"]

        lines = [
            ("1. التدفقات النقدية من الأنشطة التشغيلية:", None, True, "SECTION"),
            ("  صافي الربح قبل الزكاة", op["profit_before_zakat"], True, None),
            ("  تسويات البنود غير النقدية:", None, True, None),
            ("    + مصروف إهلاك الأصول الثابتة", op["depreciation_add_back"], False, None),
            ("    + مصروف مخصص مكافأة نهاية الخدمة", op["eos_provision_add_back"], False, None),
            ("  التغيرات في رأس المال العامل:", None, True, None),
            ("    - التغير في العملاء والمدينين التجاريين", op["change_in_receivables"], False, None),
            ("    - التغير في المخزون السلعي", op["change_in_inventory"], False, None),
            ("    - التغير في المدفوعات المقدمة وضريبة المدخلات", op["change_in_prepayments_and_input_vat"], False, None),
            ("    + التغير في الموردين والدائنين التجاريين", op["change_in_payables"], False, None),
            ("    + التغير في المصروفات المستحقة وضريبة المخرجات", op["change_in_accruals_and_output_vat"], False, None),
            ("صافي النقد الناتج من (المستخدم في) الأنشطة التشغيلية", op["net_cash_from_operating"], True, "TOTAL"),

            ("2. التدفقات النقدية من الأنشطة الاستثمارية:", None, True, "SECTION"),
            ("  - اقتناء وشراء أصول ثابتة ملموسة", inv["purchase_of_ppe"], False, None),
            ("  - اقتناء أصول غير ملموسة وبرمجيات", inv["purchase_of_intangibles"], False, None),
            ("صافي النقد المستخدم في الأنشطة الاستثمارية", inv["net_cash_used_in_investing"], True, "TOTAL"),

            ("3. التدفقات النقدية من الأنشطة التمويلية:", None, True, "SECTION"),
            ("  + المتحصلات من إصدار رأس المال والاحتياطيات", fin["capital_and_reserves_proceeds"], False, None),
            ("  + صافي حركة جاري الشركاء / المالك", fin["partners_current_movement"], False, None),
            ("  + صافي القروض والتسهيلات طويلة الأجل", fin["long_term_loans_proceeds"], False, None),
            ("صافي النقد الناتج من الأنشطة التمويلية", fin["net_cash_from_financing"], True, "TOTAL"),

            ("صافي التغير في النقد وما في حكمه خلال الفترة", cf["net_change_in_cash"], True, "GRAND_TOTAL"),
            ("💰 رصيد النقد وما في حكمه في نهاية الفترة", cf["closing_cash_balance"], True, "GRAND_TOTAL")
        ]

        cls._write_styled_statement_lines(ws, lines)

    @classmethod
    def _create_zakat_sheet(cls, wb, company_name: str, year: int, zakat: Dict[str, Any]):
        ws = wb.create_sheet(title="إقرار وعاء الزكاة (ZATCA)")
        ws.views.sheetView[0].rightToLeft = True

        ws.merge_cells("B2:D2")
        ws["B2"] = f"كشف احتساب وعاء الزكاة الشرعي التقديري طبقاً للائحة ZATCA - لعام {year}"
        ws["B2"].font = cls.FONT_TITLE
        ws["B2"].fill = cls.HEADER_FILL
        ws["B2"].alignment = cls.ALIGN_CENTER

        src = zakat["sources_added"]
        ded = zakat["deductions"]

        lines = [
            ("1. عناصر مصادر التمويل المضافة للوعاء:", None, True, "SECTION"),
            ("  + رأس المال المدفوع", src["capital"], False, None),
            ("  + الاحتياطيات النظامية والقانونية", src["reserves"], False, None),
            ("  + الأرباح المبقاة (المرحلة أول المدة)", src["retained_earnings_opening"], False, None),
            ("  + القروض والتسهيلات طويلة الأجل", src["long_term_debt"], False, None),
            ("  + مخصص مكافأة نهاية الخدمة للموظفين", src["eos_provision"], False, None),
            ("  + رصيد جاري الشركاء الدائن", src["partners_credit_balance"], False, None),
            ("إجمالي مصادر التمويل المضافة", src["total_added"], True, "TOTAL"),

            ("2. الحطوم (العناصر المحسومة من الوعاء):", None, True, "SECTION"),
            ("  - صافي الأصول الثابتة الملموسة (بعد الإهلاك)", -ded["net_fixed_assets"], False, None),
            ("  - الأصول غير الملموسة والبرمجيات", -ded["intangibles"], False, None),
            ("إجمالي الحطوم المحسومة", -ded["total_deductions"], True, "TOTAL"),

            ("3. الوعاء الزكوي وصافي الربح المعدل:", None, True, "SECTION"),
            ("الوعاء قبل صافي الربح", zakat["base_before_profit"], True, None),
            ("يضاف: صافي الربح المعدل للفترة", zakat["adjusted_net_profit"], False, None),
            ("⭐ وعاء الزكاة الشرعي الخاضع للحسبة", zakat["zakat_base"], True, "GRAND_TOTAL"),
            ("عدد أيام الفترة المالية الخاضعة", f"{zakat['total_days']} يوماً", False, None),
            ("النسبة الزكوية المطبقة (التناسبية)", f"{zakat['applicable_zakat_rate_pct']}%", False, None),
            ("💰 قيمة الزكاة الشرعية المستحقة", zakat["zakat_due"], True, "GRAND_TOTAL")
        ]

        cls._write_styled_statement_lines(ws, lines)

    @classmethod
    def _create_vat_return_sheet(cls, wb, company_name: str, year: int, vat: Dict[str, Any]):
        ws = wb.create_sheet(title="إقرار القيمة المضافة VAT")
        ws.views.sheetView[0].rightToLeft = True

        ws.merge_cells("B2:D2")
        ws["B2"] = f"نموذج إقرار ضريبة القيمة المضافة (ZATCA VAT Return) - لعام {year}"
        ws["B2"].font = cls.FONT_TITLE
        ws["B2"].fill = cls.HEADER_FILL
        ws["B2"].alignment = cls.ALIGN_CENTER

        s = vat["sales_section"]
        p = vat["purchases_section"]

        lines = [
            ("أولاً: ضريبة القيمة المضافة على المبيعات (ضريبة المخرجات):", None, True, "SECTION"),
            ("  1. المبيعات الخاضعة للنسبة الأساسية (15%)", s["standard_rated_sales_15"], False, None),
            ("  2. يخصم: مردودات المبيعات وإشعارات الدائن", -s["sales_returns_and_credit_notes"], False, None),
            ("  صافي المبيعات الخاضعة للضريبة", s["net_taxable_sales"], True, "TOTAL"),
            ("  إجمالي ضريبة المخرجات على المبيعات (15%)", s["gross_output_vat"], False, None),
            ("  تسويات ضريبة المخرجات (مرتجعات المبيعات)", -s["output_vat_adjustments"], False, None),
            ("إجمالي ضريبة المخرجات المستحقة (Total Output VAT)", s["total_output_vat"], True, "TOTAL"),

            ("ثانياً: ضريبة القيمة المضافة على المشتريات والمصروفات (ضريبة المدخلات):", None, True, "SECTION"),
            ("  3. المشتريات والمصروفات الخاضعة للنسبة الأساسية (15%)", p["standard_rated_purchases_15"], False, None),
            ("  4. يخصم: مردودات المشتريات وإشعارات المدين", -p["purchase_returns_and_debit_notes"], False, None),
            ("  صافي المشتريات الخاضعة للضريبة", p["net_taxable_purchases"], True, "TOTAL"),
            ("  إجمالي ضريبة المدخلات على المشتريات والمصروفات", p["gross_input_vat"], False, None),
            ("  تسويات ضريبة المدخلات (مرتجعات المشتريات)", -p["input_vat_adjustments"], False, None),
            ("إجمالي ضريبة المدخلات القابلة للاسترداد (Total Input VAT)", p["total_recoverable_input_vat"], True, "TOTAL"),

            ("ثالثاً: صافي الضريبة المستحقة أو القابلة للاسترداد:", None, True, "SECTION"),
            ("🏛️ صافي ضريبة القيمة المضافة المستحقة (Net VAT Due)", vat["net_vat_due"], True, "GRAND_TOTAL"),
            ("حالة الإقرار الضريبي", vat["settlement_status"], True, None)
        ]

        cls._write_styled_statement_lines(ws, lines)

    @classmethod
    def _create_ratios_sheet(cls, wb, company_name: str, year: int, ratios: Dict[str, Any]):
        ws = wb.create_sheet(title="المؤشرات والنسب المالية")
        ws.views.sheetView[0].rightToLeft = True

        ws.merge_cells("B2:D2")
        ws["B2"] = f"لوحة التحليل والنسب المالية التنفيذية - لعام {year}"
        ws["B2"].font = cls.FONT_TITLE
        ws["B2"].fill = cls.HEADER_FILL
        ws["B2"].alignment = cls.ALIGN_CENTER

        lines = [
            ("1. نسب السيولة ورأس المال العامل (Liquidity):", None, True, "SECTION"),
            ("نسبة التداول (Current Ratio)", f"{ratios['liquidity']['current_ratio']}x", False, None),
            ("نسبة السيولة السريعة (Quick Ratio)", f"{ratios['liquidity']['quick_ratio']}x", False, None),
            ("رأس المال العامل الصافي (Working Capital)", ratios["liquidity"]["working_capital"], True, "TOTAL"),

            ("2. نسب الربحية والأداء المالي (Profitability):", None, True, "SECTION"),
            ("هامش مجمل الربح (Gross Profit Margin)", f"{ratios['profitability']['gross_margin_pct']}%", False, None),
            ("هامش الربح التشغيلي (Operating Margin)", f"{ratios['profitability']['operating_margin_pct']}%", False, None),
            ("هامش صافي الربح (Net Profit Margin)", f"{ratios['profitability']['net_margin_pct']}%", False, None),
            ("العائد على إجمالي الأصول (ROA)", f"{ratios['profitability']['return_on_assets_pct']}%", False, None),
            ("العائد على حقوق الملكية (ROE)", f"{ratios['profitability']['return_on_equity_pct']}%", False, None),

            ("3. نسب الملاءة والرافعة المالية (Solvency):", None, True, "SECTION"),
            ("نسبة الديون إلى حقوق الملكية (Debt to Equity)", f"{ratios['solvency']['debt_to_equity_ratio']}x", False, None),
            ("نسبة الديون إلى الأصول (Debt to Assets)", f"{ratios['solvency']['debt_to_assets_ratio']}x", False, None)
        ]

        cls._write_styled_statement_lines(ws, lines)

    @classmethod
    def _create_dataframe_sheet(cls, wb, sheet_title: str, df: pd.DataFrame):
        ws = wb.create_sheet(title=sheet_title)
        ws.views.sheetView[0].rightToLeft = True

        headers = list(df.columns) if not df.empty else ["البيان", "الحالة"]
        for col_num, h_text in enumerate(headers, start=1):
            cell = ws.cell(row=2, column=col_num, value=h_text)
            cell.font = cls.FONT_HEADER
            cell.fill = cls.HEADER_FILL
            cell.alignment = cls.ALIGN_CENTER
            cell.border = cls.THIN_BORDER

        row_idx = 3
        if df.empty:
            ws.cell(row=3, column=1, value="لا توجد سجلات حالية").alignment = cls.ALIGN_CENTER
        else:
            for _, r in df.iterrows():
                for c_idx, h in enumerate(headers, start=1):
                    val = r[h]
                    c_cell = ws.cell(row=row_idx, column=c_idx, value=val)
                    c_cell.border = cls.THIN_BORDER
                    c_cell.font = cls.FONT_REGULAR
                    if isinstance(val, (int, float)):
                        c_cell.number_format = "#,##0.00"
                        c_cell.alignment = cls.ALIGN_LEFT
                    else:
                        c_cell.alignment = cls.ALIGN_RIGHT if c_idx == 2 else cls.ALIGN_CENTER
                row_idx += 1

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

    @classmethod
    def _create_notes_sheet(cls, wb, company_name: str, year: int):
        ws = wb.create_sheet(title="الإيضاحات والسياسات المحاسبية")
        ws.views.sheetView[0].rightToLeft = True

        ws.merge_cells("B2:D2")
        ws["B2"] = "ملخص أهم السياسات والإيضاحات المحاسبية المتممة للقوائم المالية"
        ws["B2"].font = cls.FONT_TITLE
        ws["B2"].fill = cls.HEADER_FILL
        ws["B2"].alignment = cls.ALIGN_CENTER

        policies = [
            ("1. أساس إعداد القوائم المالية", "تم إعداد القوائم المالية وفقاً للمعايير الدولية للتقرير المالي للمنشآت الصغيرة والمتوسطة (IFRS for SMEs) المعتمدة من الهيئة السعودية للمراجعين والمحاسبين (SOCPA) وعلى أساس الاستحقاق ومبدأ الاستمرارية."),
            ("2. الاعتراف بالإيرادات", "يتم الاعتراف بإيرادات المبيعات عند انتقال السيطرة على البضائع أو تسليم الخدمات للعملاء وفقاً لمعيار IFRS 15 بصافي المبالغ بعد استبعاد الخصومات ومردودات المبيعات وضريبة القيمة المضافة."),
            ("3. تقييم المخزون السلعي", "يُقيّم المخزون السلعي بالتكلفة أو صافي القيمة القابلة للتحقق أيهما أقل، وتُحدد التكلفة آلياً باستخدام طريقة المتوسط المرجح المتحرك (Moving Weighted Average Cost) بعد كل عملية توريد."),
            ("4. الأصول الثابتة وإهلاكها", "تُثبت الأصول الثابتة بالتكلفة التاريخية ناقصاً مجمع الإهلاك وخسائر الهبوط، ويُحسب الإهلاك بطريقة القسط الثابت على مدار الأعمار الإنتاجية المقدرة للأصول."),
            ("5. احتساب الزكاة الشرعية", "تُحسب الزكاة الشرعية التقديرية وفقاً للائحة التنفيذية لجباية الزكاة الصادرة عن هيئة الزكاة والضريبة والجمارك (ZATCA) في المملكة العربية السعودية على أساس وعاء الزكاة المعدل أو الربح المعدل أيهما أكبر."),
            ("6. ضريبة القيمة المضافة (VAT)", "تُطبق المنشأة اللائحة التنفيذية لضريبة القيمة المضافة بنسبة 15% وتصدر الفواتير الضريبية وإشعارات الدائن والمدين المشفرة برمز الاستجابة السريعة (ZATCA TLV QR Code).")
        ]

        curr_row = 4
        for title, desc in policies:
            ws.cell(row=curr_row, column=2, value=title).font = cls.FONT_BOLD
            ws.cell(row=curr_row, column=2).fill = cls.SECTION_FILL
            ws.merge_cells(start_row=curr_row, start_column=2, end_row=curr_row, end_column=4)
            curr_row += 1

            d_cell = ws.cell(row=curr_row, column=2, value=desc)
            d_cell.font = cls.FONT_REGULAR
            d_cell.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
            ws.merge_cells(start_row=curr_row, start_column=2, end_row=curr_row, end_column=4)
            ws.row_dimensions[curr_row].height = 45
            curr_row += 2

        ws.column_dimensions["B"].width = 30
        ws.column_dimensions["C"].width = 30
        ws.column_dimensions["D"].width = 30

    @classmethod
    def _create_journal_log_sheet(cls, wb, db: Session, start_date: Optional[date] = None, end_date: Optional[date] = None):
        ws = wb.create_sheet(title="دفتر اليومية العامة")
        ws.views.sheetView[0].rightToLeft = True

        headers = ["رقم القيد", "التاريخ", "نوع المستند", "المرجع", "بيان القيد", "كود الحساب", "اسم الحساب", "مدين", "دائن", "حالة القيد"]

        for col_num, h_text in enumerate(headers, start=1):
            cell = ws.cell(row=2, column=col_num, value=h_text)
            cell.font = cls.FONT_HEADER
            cell.fill = cls.HEADER_FILL
            cell.alignment = cls.ALIGN_CENTER
            cell.border = cls.THIN_BORDER

        from models import EntryStatus
        q_entries = db.query(JournalEntry).filter(JournalEntry.status != EntryStatus.DRAFT)
        if start_date:
            q_entries = q_entries.filter(JournalEntry.entry_date >= start_date)
        if end_date:
            q_entries = q_entries.filter(JournalEntry.entry_date <= end_date)
        entries = q_entries.order_by(JournalEntry.entry_date, JournalEntry.id).all()
        row_idx = 3

        for entry in entries:
            for line in entry.lines:
                ws.cell(row=row_idx, column=1, value=entry.entry_number).alignment = cls.ALIGN_CENTER
                ws.cell(row=row_idx, column=2, value=entry.entry_date.strftime("%Y-%m-%d")).alignment = cls.ALIGN_CENTER
                ws.cell(row=row_idx, column=3, value=entry.doc_type.value).alignment = cls.ALIGN_CENTER
                ws.cell(row=row_idx, column=4, value=entry.reference or "").alignment = cls.ALIGN_CENTER
                ws.cell(row=row_idx, column=5, value=line.description or entry.description).alignment = cls.ALIGN_RIGHT
                ws.cell(row=row_idx, column=6, value=line.account.code).alignment = cls.ALIGN_CENTER
                ws.cell(row=row_idx, column=7, value=line.account.name_ar).alignment = cls.ALIGN_RIGHT

                deb_cell = ws.cell(row=row_idx, column=8, value=line.debit)
                deb_cell.number_format = "#,##0.00;(#,##0.00);\"-\""
                deb_cell.alignment = cls.ALIGN_LEFT

                crd_cell = ws.cell(row=row_idx, column=9, value=line.credit)
                crd_cell.number_format = "#,##0.00;(#,##0.00);\"-\""
                crd_cell.alignment = cls.ALIGN_LEFT

                ws.cell(row=row_idx, column=10, value=entry.status.value).alignment = cls.ALIGN_CENTER
                for col in range(1, 11):
                    ws.cell(row=row_idx, column=col).border = cls.THIN_BORDER
                    ws.cell(row=row_idx, column=col).font = cls.FONT_REGULAR

                row_idx += 1

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 13)

    @classmethod
    def _write_styled_statement_lines(cls, ws, lines):
        curr_row = 4
        for title, val, is_bold, style_type in lines:
            title_cell = ws.cell(row=curr_row, column=2, value=title)
            title_cell.font = cls.FONT_BOLD if is_bold else cls.FONT_REGULAR
            title_cell.alignment = cls.ALIGN_RIGHT
            title_cell.border = cls.THIN_BORDER

            val_cell = ws.cell(row=curr_row, column=3, value=val if val is not None else "")
            val_cell.font = cls.FONT_BOLD if is_bold else cls.FONT_REGULAR
            val_cell.alignment = cls.ALIGN_LEFT
            val_cell.border = cls.THIN_BORDER

            if isinstance(val, (int, float)):
                val_cell.number_format = "#,##0.00;(#,##0.00);\"-\""

            if style_type == "SECTION":
                title_cell.fill = cls.SUBHEADER_FILL
                title_cell.font = cls.FONT_HEADER
                ws.merge_cells(start_row=curr_row, start_column=2, end_row=curr_row, end_column=3)
            elif style_type == "TOTAL":
                title_cell.fill = cls.SECTION_FILL
                val_cell.fill = cls.SECTION_FILL
            elif style_type == "GRAND_TOTAL":
                title_cell.fill = cls.TOTAL_FILL
                val_cell.fill = cls.TOTAL_FILL
                val_cell.border = cls.DOUBLE_BOTTOM_BORDER

            curr_row += 1

        ws.column_dimensions["B"].width = 50
        ws.column_dimensions["C"].width = 25


# =============================================================================
# 2. محرك قوالب الطباعة الرسمية الموحدة (HTML Printable Templates - Offline Ready)
# =============================================================================

class PrintTemplateService:
    """توليد قوالب الطباعة الرسمية للفواتير والإشعارات الضريبية والسندات ومسير الرواتب والتسوية البنكية"""

    @staticmethod
    def generate_sales_invoice_html(
        invoice: SalesInvoice,
        settings: Optional[CompanySettings] = None
    ) -> str:
        seller_name = settings.company_name_ar if settings else "شركة ميزان برو التجارية"
        seller_vat = settings.vat_number if settings else "300000000000003"
        seller_cr = settings.cr_number if settings else "1010889922"
        seller_addr = settings.address_ar if settings else "المملكة العربية السعودية - الرياض"

        items_rows_html = ""
        for idx, it in enumerate(invoice.items, start=1):
            items_rows_html += f"""
            <tr>
                <td style="text-align:center; padding:8px; border:1px solid #CBD5E1;">{idx}</td>
                <td style="text-align:right; padding:8px; border:1px solid #CBD5E1;">{html.escape(str(it.item_name or ""))}</td>
                <td style="text-align:center; padding:8px; border:1px solid #CBD5E1;">{it.quantity:,.2f}</td>
                <td style="text-align:left; padding:8px; border:1px solid #CBD5E1;">{it.unit_price:,.2f}</td>
                <td style="text-align:left; padding:8px; border:1px solid #CBD5E1;">{it.discount:,.2f}</td>
                <td style="text-align:left; padding:8px; border:1px solid #CBD5E1;">{it.taxable_amount:,.2f}</td>
                <td style="text-align:center; padding:8px; border:1px solid #CBD5E1;">{int(it.tax_rate * 100)}%</td>
                <td style="text-align:left; padding:8px; border:1px solid #CBD5E1;">{it.tax_amount:,.2f}</td>
                <td style="text-align:left; padding:8px; border:1px solid #CBD5E1; font-weight:bold;">{it.total_with_tax:,.2f}</td>
            </tr>
            """

        z_rep = ZATCAPhase2Validator.validate_financial_document(invoice, settings=settings)
        qr_svg_uri = LedgerService.generate_qr_svg_data_uri(z_rep["qr_code_base64"], size=130)
        inv_type_title = "فاتورة ضريبية مبسطة (B2C - 0200000)" if invoice.invoice_type == InvoiceType.SIMPLIFIED else "فاتورة ضريبية قياسية (B2B - 0100000)"

        return f"""
        <!DOCTYPE html>
        <html lang="ar" dir="rtl">
        <head>
            <meta charset="UTF-8">
            <title>فاتورة مبيعات ضريبية - {invoice.invoice_number}</title>
            <style>
                * {{ box-sizing: border-box; font-family: 'Cairo', 'Segoe UI', Tahoma, sans-serif; }}
                body {{ margin: 0; padding: 20px; background-color: #F8FAFC; color: #0F172A; }}
                .invoice-container {{ max-width: 820px; margin: 0 auto; background: white; padding: 30px; border-radius: 12px; border: 1px solid #CBD5E1; }}
                .invoice-header {{ display: flex; justify-content: space-between; align-items: center; border-bottom: 2px solid #1E3A8A; padding-bottom: 18px; margin-bottom: 18px; }}
                .company-info h2 {{ margin: 0 0 5px 0; color: #1E3A8A; font-size: 22px; }}
                .company-info p {{ margin: 2px 0; font-size: 13px; color: #475569; }}
                .badge-title {{ background-color: #1E3A8A; color: white; padding: 6px 16px; border-radius: 6px; font-size: 15px; font-weight: bold; display: inline-block; margin-bottom: 8px; }}
                .info-grid {{ display: grid; grid-template-columns: 1fr 1fr; gap: 18px; background-color: #F1F5F9; padding: 14px; border-radius: 8px; margin-bottom: 18px; font-size: 13px; }}
                table {{ width: 100%; border-collapse: collapse; margin-bottom: 18px; font-size: 12.5px; }}
                th {{ background-color: #1E3A8A; color: white; padding: 9px; border: 1px solid #1E3A8A; text-align: center; }}
                .totals-container {{ display: flex; justify-content: space-between; align-items: flex-end; margin-top: 18px; padding-top: 15px; border-top: 2px solid #E2E8F0; }}
                .totals-box {{ width: 340px; background: #F8FAFC; border: 1px solid #CBD5E1; border-radius: 8px; padding: 12px; font-size: 13px; }}
                .totals-row {{ display: flex; justify-content: space-between; margin-bottom: 6px; }}
                .grand-total {{ border-top: 2px solid #1E3A8A; padding-top: 8px; font-size: 16px; font-weight: bold; color: #1E3A8A; }}
                .print-btn {{ background-color: #16A34A; color: white; border: none; padding: 11px 22px; border-radius: 6px; font-size: 14px; font-weight: bold; cursor: pointer; display: block; width: 100%; margin-bottom: 18px; }}
                @media print {{ body {{ background: white; padding: 0; }} .invoice-container {{ border: none; padding: 10px; max-width: 100%; }} .no-print {{ display: none !important; }} }}
            </style>
        </head>
        <body>
            <div class="invoice-container">
                <button class="print-btn no-print" onclick="window.print()">🖨️ طباعة الفاتورة الضريبية الآن (أو الحفظ كـ PDF)</button>
                <div class="invoice-header">
                    <div class="company-info">
                        <h2>{seller_name}</h2>
                        <p><b>الرقم الضريبي للمنشأة:</b> {seller_vat}</p>
                        <p><b>السجل التجاري:</b> {seller_cr} | <b>العنوان:</b> {seller_addr}</p>
                    </div>
                    <div style="text-align:left;">
                        <div class="badge-title">{inv_type_title}</div>
                        <p style="margin:0; font-size:12px; color:#64748B;">ZATCA E-Invoice Standard</p>
                    </div>
                </div>
                <div class="info-grid">
                    <div>
                        <p><b>رقم الفاتورة:</b> <span style="font-family:monospace; font-weight:bold;">{invoice.invoice_number}</span></p>
                        <p><b>تاريخ الإصدار:</b> {invoice.invoice_date}</p>
                        <p><b>طريقة السداد:</b> {invoice.payment_method.value}</p>
                    </div>
                    <div>
                        <p><b>اسم العميل / المشتري:</b> {html.escape(str(invoice.customer_name or ""))}</p>
                        <p><b>الرقم الضريبي للعميل:</b> {html.escape(str(invoice.customer_vat_number or "غير متوفر (عميل فرد)"))}</p>
                        <p><b>الحالة:</b> معتمدة ومرحلة بالدفاتر</p>
                    </div>
                </div>
                <table>
                    <thead>
                        <tr>
                            <th>#</th><th>اسم الصنف / الخدمة</th><th>الكمية</th><th>السعر</th><th>الخصم</th><th>الخاضع للضريبة</th><th>النسبة</th><th>الضريبة</th><th>الإجمالي</th>
                        </tr>
                    </thead>
                    <tbody>{items_rows_html}</tbody>
                </table>
                <div class="totals-container">
                    <div style="text-align:center;">
                        <img src="{qr_svg_uri}" width="125" height="125" style="border:1px solid #CBD5E1; padding:4px; border-radius:6px; background:white;"/>
                        <p style="font-size:11px; color:#0F172A; font-weight:bold; margin:4px 0 0 0;">ZATCA Phase 2 QR (TLV Tags 1-9) ✅</p>
                        <p style="font-size:9.5px; color:#64748B; font-family:monospace; margin:2px 0 0 0;">UUID: {z_rep["uuid"][:18]}...</p>
                    </div>
                    <div class="totals-box">
                        <div class="totals-row"><span>المجموع قبل الخصم:</span><b>{invoice.subtotal:,.2f} ريال</b></div>
                        <div class="totals-row"><span>إجمالي الخصم:</span><b>{invoice.discount_amount:,.2f} ريال</b></div>
                        <div class="totals-row"><span>الصافي الخاضع للضريبة:</span><b>{invoice.taxable_amount:,.2f} ريال</b></div>
                        <div class="totals-row"><span>ضريبة القيمة المضافة (15%):</span><b>{invoice.vat_amount:,.2f} ريال</b></div>
                        <div class="totals-row grand-total"><span>الإجمالي النهائي المستحق:</span><span>{invoice.grand_total:,.2f} ريال</span></div>
                    </div>
                </div>
                {f'<div style="margin-top:15px; font-size:12px; color:#64748B;"><b>ملاحظات:</b> {html.escape(str(invoice.notes or ""))}</div>' if invoice.notes else ''}
            </div>
        </body>
        </html>
        """

    @staticmethod
    def generate_purchase_invoice_html(
        invoice: PurchaseInvoice,
        settings: Optional[CompanySettings] = None
    ) -> str:
        company_name = settings.company_name_ar if settings else "شركة ميزان برو التجارية"
        company_vat = settings.vat_number if settings else "300000000000003"

        items_rows_html = ""
        for idx, it in enumerate(invoice.items, start=1):
            items_rows_html += f"""
            <tr>
                <td style="text-align:center; padding:8px; border:1px solid #CBD5E1;">{idx}</td>
                <td style="text-align:right; padding:8px; border:1px solid #CBD5E1;">{html.escape(str(it.item_name or ""))}</td>
                <td style="text-align:center; padding:8px; border:1px solid #CBD5E1;">{it.quantity:,.2f}</td>
                <td style="text-align:left; padding:8px; border:1px solid #CBD5E1;">{it.unit_price:,.2f}</td>
                <td style="text-align:left; padding:8px; border:1px solid #CBD5E1;">{it.discount:,.2f}</td>
                <td style="text-align:left; padding:8px; border:1px solid #CBD5E1;">{it.taxable_amount:,.2f}</td>
                <td style="text-align:left; padding:8px; border:1px solid #CBD5E1;">{it.tax_amount:,.2f}</td>
                <td style="text-align:left; padding:8px; border:1px solid #CBD5E1; font-weight:bold;">{it.total_with_tax:,.2f}</td>
            </tr>
            """

        return f"""
        <!DOCTYPE html>
        <html lang="ar" dir="rtl">
        <head>
            <meta charset="UTF-8">
            <title>فاتورة مشتريات وسند استلام مخزني - {invoice.invoice_number}</title>
            <style>
                * {{ box-sizing: border-box; font-family: 'Cairo', 'Segoe UI', Tahoma, sans-serif; }}
                body {{ margin: 0; padding: 20px; background-color: #F8FAFC; color: #0F172A; }}
                .invoice-container {{ max-width: 820px; margin: 0 auto; background: white; padding: 30px; border-radius: 12px; border: 1px solid #CBD5E1; }}
                .invoice-header {{ display: flex; justify-content: space-between; align-items: center; border-bottom: 2px solid #0F766E; padding-bottom: 18px; margin-bottom: 18px; }}
                .company-info h2 {{ margin: 0 0 5px 0; color: #0F766E; font-size: 22px; }}
                .company-info p {{ margin: 2px 0; font-size: 13px; color: #475569; }}
                .badge-title {{ background-color: #0F766E; color: white; padding: 6px 16px; border-radius: 6px; font-size: 15px; font-weight: bold; display: inline-block; margin-bottom: 8px; }}
                .info-grid {{ display: grid; grid-template-columns: 1fr 1fr; gap: 18px; background-color: #F0FDFA; padding: 14px; border-radius: 8px; margin-bottom: 18px; font-size: 13px; }}
                table {{ width: 100%; border-collapse: collapse; margin-bottom: 18px; font-size: 12.5px; }}
                th {{ background-color: #0F766E; color: white; padding: 9px; border: 1px solid #0F766E; text-align: center; }}
                .totals-container {{ display: flex; justify-content: flex-end; margin-top: 18px; padding-top: 15px; border-top: 2px solid #E2E8F0; }}
                .totals-box {{ width: 340px; background: #F8FAFC; border: 1px solid #CBD5E1; border-radius: 8px; padding: 12px; font-size: 13px; }}
                .totals-row {{ display: flex; justify-content: space-between; margin-bottom: 6px; }}
                .grand-total {{ border-top: 2px solid #0F766E; padding-top: 8px; font-size: 16px; font-weight: bold; color: #0F766E; }}
                .print-btn {{ background-color: #0F766E; color: white; border: none; padding: 11px 22px; border-radius: 6px; font-size: 14px; font-weight: bold; cursor: pointer; display: block; width: 100%; margin-bottom: 18px; }}
                @media print {{ body {{ background: white; padding: 0; }} .invoice-container {{ border: none; padding: 10px; max-width: 100%; }} .no-print {{ display: none !important; }} }}
            </style>
        </head>
        <body>
            <div class="invoice-container">
                <button class="print-btn no-print" onclick="window.print()">🖨️ طباعة سند المشتريات والاستلام المخزني (أو الحفظ كـ PDF)</button>
                <div class="invoice-header">
                    <div class="company-info">
                        <h2>{company_name}</h2>
                        <p><b>الرقم الضريبي للمشتري:</b> {company_vat}</p>
                        <p><b>القسم:</b> المشتريات وإدارة المستودعات (تقييم المتوسط المرجح WAC)</p>
                    </div>
                    <div style="text-align:left;">
                        <div class="badge-title">فاتورة مشتريات / إذن استلام</div>
                        <p style="margin:0; font-size:12px; color:#64748B;">Purchase Invoice & GRN</p>
                    </div>
                </div>
                <div class="info-grid">
                    <div>
                        <p><b>رقم سند الشراء:</b> <span style="font-family:monospace; font-weight:bold;">{invoice.invoice_number}</span></p>
                        <p><b>رقم فاتورة المورد:</b> {invoice.supplier_invoice_number or "غير محدد"}</p>
                        <p><b>تاريخ الاستلام:</b> {invoice.invoice_date}</p>
                    </div>
                    <div>
                        <p><b>اسم المورد:</b> {html.escape(str(invoice.supplier_name or ""))}</p>
                        <p><b>الرقم الضريبي للمورد:</b> {invoice.supplier_vat_number or "غير متوفر"}</p>
                        <p><b>طريقة السداد:</b> {invoice.payment_method.value}</p>
                    </div>
                </div>
                <table>
                    <thead>
                        <tr>
                            <th>#</th><th>الصنف / البضاعة المشتراة</th><th>الكمية</th><th>سعر الوحدة</th><th>الخصم</th><th>الصافي الخاضع</th><th>مبلغ الضريبة</th><th>الإجمالي</th>
                        </tr>
                    </thead>
                    <tbody>{items_rows_html}</tbody>
                </table>
                <div class="totals-container" style="justify-content:space-between; align-items:center;">
                    <div style="text-align:center;">
                        <img src="{LedgerService.generate_qr_svg_data_uri(ZATCAPhase2Validator.validate_financial_document(invoice, settings=settings)['qr_code_base64'], size=115)}" width="115" height="115" style="border:1px solid #CBD5E1; padding:4px; border-radius:6px; background:white;"/>
                        <p style="font-size:10.5px; color:#0F766E; font-weight:bold; margin:4px 0 0 0;">ZATCA Phase 2 Verified ✅</p>
                    </div>
                    <div class="totals-box">
                        <div class="totals-row"><span>صافي المشتريات (غير شامل الضريبة):</span><b>{invoice.taxable_amount:,.2f} ريال</b></div>
                        <div class="totals-row"><span>ضريبة المدخلات 15% (مستردة):</span><b>{invoice.vat_amount:,.2f} ريال</b></div>
                        <div class="totals-row grand-total"><span>إجمالي المستحق للمورد:</span><span>{invoice.grand_total:,.2f} ريال</span></div>
                    </div>
                </div>
                {f'<div style="margin-top:15px; font-size:12px; color:#64748B;"><b>ملاحظات:</b> {html.escape(str(invoice.notes or ""))}</div>' if invoice.notes else ''}
            </div>
        </body>
        </html>
        """

    @staticmethod
    def generate_credit_debit_note_html(
        note: CreditDebitNote,
        settings: Optional[CompanySettings] = None
    ) -> str:
        company_name = settings.company_name_ar if settings else "شركة ميزان برو التجارية"
        company_vat = settings.vat_number if settings else "300000000000003"
        is_credit = note.note_type == NoteType.CREDIT_NOTE
        theme_color = "#B91C1C" if is_credit else "#7C3AED"
        title_ar = "إشعار دائن ضريبي (مرتجع مبيعات)" if is_credit else "إشعار مدين ضريبي (مرتجع مشتريات)"
        party_label = "العميل" if is_credit else "المورد"

        items_rows = ""
        for idx, it in enumerate(note.items, start=1):
            items_rows += f"""
            <tr>
                <td style="text-align:center; padding:8px; border:1px solid #CBD5E1;">{idx}</td>
                <td style="text-align:right; padding:8px; border:1px solid #CBD5E1;">{html.escape(str(it.item_name or ""))}</td>
                <td style="text-align:center; padding:8px; border:1px solid #CBD5E1;">{it.quantity:,.2f}</td>
                <td style="text-align:left; padding:8px; border:1px solid #CBD5E1;">{it.unit_price:,.2f}</td>
                <td style="text-align:left; padding:8px; border:1px solid #CBD5E1;">{it.taxable_amount:,.2f}</td>
                <td style="text-align:left; padding:8px; border:1px solid #CBD5E1;">{it.tax_amount:,.2f}</td>
                <td style="text-align:left; padding:8px; border:1px solid #CBD5E1; font-weight:bold;">{it.total_with_tax:,.2f}</td>
            </tr>
            """

        qr_svg_uri = LedgerService.generate_qr_svg_data_uri(note.qr_code_base64 or note.note_number, size=125)

        return f"""
        <!DOCTYPE html>
        <html lang="ar" dir="rtl">
        <head>
            <meta charset="UTF-8">
            <title>{title_ar} - {note.note_number}</title>
            <style>
                * {{ box-sizing: border-box; font-family: 'Cairo', 'Segoe UI', Tahoma, sans-serif; }}
                body {{ margin: 0; padding: 20px; background-color: #F8FAFC; color: #0F172A; }}
                .doc-box {{ max-width: 820px; margin: 0 auto; background: white; padding: 28px; border-radius: 12px; border: 1px solid #CBD5E1; }}
                .doc-header {{ display: flex; justify-content: space-between; align-items: center; border-bottom: 2px solid {theme_color}; padding-bottom: 16px; margin-bottom: 16px; }}
                .badge {{ background-color: {theme_color}; color: white; padding: 6px 16px; border-radius: 6px; font-weight: bold; font-size: 15px; }}
                .info-grid {{ display: grid; grid-template-columns: 1fr 1fr; gap: 16px; background: #F8FAFC; padding: 14px; border-radius: 8px; border: 1px solid #E2E8F0; margin-bottom: 16px; font-size: 13px; }}
                table {{ width: 100%; border-collapse: collapse; margin-bottom: 16px; font-size: 12.5px; }}
                th {{ background-color: {theme_color}; color: white; padding: 9px; border: 1px solid {theme_color}; }}
                .print-btn {{ background-color: {theme_color}; color: white; border: none; padding: 10px 20px; border-radius: 6px; font-weight: bold; cursor: pointer; width: 100%; margin-bottom: 16px; }}
                @media print {{ .no-print {{ display: none !important; }} body {{ background: white; padding: 0; }} .doc-box {{ border: none; }} }}
            </style>
        </head>
        <body>
            <div class="doc-box">
                <button class="print-btn no-print" onclick="window.print()">🖨️ طباعة الإشعار الضريبي الآن</button>
                <div class="doc-header">
                    <div>
                        <h2 style="margin:0 0 4px 0; color:{theme_color};">{company_name}</h2>
                        <p style="margin:0; font-size:13px;">الرقم الضريبي: {company_vat}</p>
                    </div>
                    <div style="text-align:left;">
                        <div class="badge">{title_ar}</div>
                        <p style="margin:4px 0 0 0; font-size:12px; color:#64748B;">ZATCA Compliant Note</p>
                    </div>
                </div>
                <div class="info-grid">
                    <div>
                        <p><b>رقم الإشعار:</b> {note.note_number}</p>
                        <p><b>تاريخ الإشعار:</b> {note.note_date}</p>
                        <p><b>الفاتورة الأصلية المرتبطة:</b> {note.original_invoice_number or "غير محدد"}</p>
                    </div>
                    <div>
                        <p><b>اسم {party_label}:</b> {html.escape(str(note.party_name or ""))}</p>
                        <p><b>الرقم الضريبي لـ{party_label}:</b> {note.party_vat_number or "-"}</p>
                        <p><b>سبب الإصدار:</b> {html.escape(str(note.reason or ""))}</p>
                    </div>
                </div>
                <table>
                    <thead>
                        <tr>
                            <th>#</th><th>الصنف</th><th>الكمية المرتجعة</th><th>سعر الوحدة</th><th>الخاضع للضريبة</th><th>الضريبة (15%)</th><th>الإجمالي</th>
                        </tr>
                    </thead>
                    <tbody>{items_rows}</tbody>
                </table>
                <div style="display:flex; justify-content:space-between; align-items:center; border-top:2px solid #E2E8F0; padding-top:14px;">
                    <div style="text-align:center;">
                        <img src="{qr_svg_uri}" width="115" height="115" style="border:1px solid #CBD5E1; padding:4px; border-radius:6px;"/>
                    </div>
                    <div style="width:320px; background:#F8FAFC; padding:12px; border-radius:8px; border:1px solid #CBD5E1; font-size:13px;">
                        <div style="display:flex; justify-content:space-between; margin-bottom:6px;"><span>القيمة قبل الضريبة:</span><b>{note.taxable_amount:,.2f} ريال</b></div>
                        <div style="display:flex; justify-content:space-between; margin-bottom:6px;"><span>ضريبة القيمة المضافة (15%):</span><b>{note.vat_amount:,.2f} ريال</b></div>
                        <div style="display:flex; justify-content:space-between; border-top:2px solid {theme_color}; padding-top:6px; font-weight:bold; font-size:15px; color:{theme_color};"><span>إجمالي قيمة الإشعار:</span><span>{note.grand_total:,.2f} ريال</span></div>
                    </div>
                </div>
            </div>
        </body>
        </html>
        """

    @staticmethod
    def generate_cash_voucher_html(
        voucher: CashVoucher,
        settings: Optional[CompanySettings] = None
    ) -> str:
        company_name = settings.company_name_ar if settings else "شركة ميزان برو التجارية"
        company_vat = settings.vat_number if settings else "300000000000003"
        is_receipt = voucher.voucher_type == VoucherType.RECEIPT
        color = "#15803D" if is_receipt else "#C2410C"
        title_ar = "سند قبض نقدية / بنكي (Receipt Voucher)" if is_receipt else "سند صرف نقدية / بنكي (Payment Voucher)"
        party_title = "استلمنا من السادة / المكرم:" if is_receipt else "صُرف للسادة / المكرم:"

        return f"""
        <!DOCTYPE html>
        <html lang="ar" dir="rtl">
        <head>
            <meta charset="UTF-8">
            <title>{title_ar} - {voucher.voucher_number}</title>
            <style>
                * {{ box-sizing: border-box; font-family: 'Cairo', 'Segoe UI', Tahoma, sans-serif; }}
                body {{ margin: 0; padding: 20px; background-color: #F8FAFC; color: #0F172A; }}
                .v-box {{ max-width: 760px; margin: 0 auto; background: white; padding: 28px; border-radius: 12px; border: 2px solid {color}; }}
                .v-header {{ display: flex; justify-content: space-between; align-items: center; border-bottom: 2px solid {color}; padding-bottom: 14px; margin-bottom: 18px; }}
                .v-badge {{ background: {color}; color: white; padding: 8px 18px; border-radius: 6px; font-weight: bold; font-size: 16px; }}
                .v-row {{ display: flex; justify-content: space-between; padding: 10px 14px; background: #F8FAFC; border: 1px solid #E2E8F0; border-radius: 6px; margin-bottom: 10px; font-size: 14px; }}
                .print-btn {{ background: {color}; color: white; border: none; padding: 10px; width: 100%; border-radius: 6px; font-weight: bold; cursor: pointer; margin-bottom: 15px; }}
                @media print {{ .no-print {{ display: none !important; }} body {{ background: white; padding: 0; }} }}
            </style>
        </head>
        <body>
            <div class="v-box">
                <button class="print-btn no-print" onclick="window.print()">🖨️ طباعة السند الرسمي الآن</button>
                <div class="v-header">
                    <div>
                        <h2 style="margin:0 0 4px 0; color:{color};">{company_name}</h2>
                        <p style="margin:0; font-size:12px; color:#475569;">الرقم الضريبي: {company_vat}</p>
                    </div>
                    <div class="v-badge">{title_ar}</div>
                </div>
                <div class="v-row">
                    <span><b>رقم السند:</b> {voucher.voucher_number}</span>
                    <span><b>التاريخ:</b> {voucher.voucher_date}</span>
                    <span><b>طريقة الدفع:</b> {voucher.payment_method.value}</span>
                </div>
                <div class="v-row">
                    <span><b>{party_title}</b> {html.escape(str(voucher.party_name or ""))}</span>
                    <span><b>الحساب المحاسبي:</b> {voucher.target_account_code}</span>
                </div>
                <div class="v-row" style="background:#FEFCE8; border-color:#FDE047; font-size:16px;">
                    <span><b>المبلغ الإجمالي:</b> {voucher.total_amount:,.2f} ريال سعودي</span>
                    <span><b>الضريبة المتضمنة:</b> {voucher.vat_amount:,.2f} ريال</span>
                </div>
                <div class="v-row">
                    <span><b>وذلك عن (البيان):</b> {html.escape(str(voucher.description or ""))}</span>
                    <span><b>المرجع:</b> {voucher.reference_number or "-"}</span>
                </div>
                <div style="margin-top:18px; display:flex; justify-content:space-between; align-items:center; border-top:1px solid #E2E8F0; padding-top:12px;">
                    <div style="text-align:center;">
                        <img src="{LedgerService.generate_qr_svg_data_uri(ZATCAPhase2Validator.validate_financial_document(voucher, settings=settings)['qr_code_base64'], size=95)}" width="95" height="95" style="border:1px solid #CBD5E1; padding:3px; border-radius:6px; background:white;"/>
                        <p style="font-size:10px; color:#475569; margin:2px 0 0 0;">ZATCA Phase 2 TLV QR</p>
                    </div>
                    <div style="flex:1; margin-right:24px; display:flex; justify-content:space-between; font-size:13px;">
                        <div><b>أمين الصندوق / الخزينة:</b> ........................</div>
                        <div><b>المحاسب المختص:</b> ........................</div>
                        <div><b>توقيع المستلم:</b> ........................</div>
                    </div>
                </div>
            </div>
        </body>
        </html>
        """

    @staticmethod
    def generate_payroll_html(
        payroll: PayrollRun,
        employees_df: pd.DataFrame,
        settings: Optional[CompanySettings] = None
    ) -> str:
        """توليد قالب مسير الرواتب الشهري الرسمي القابل للطباعة"""
        company_name = settings.company_name_ar if settings else "شركة ميزان برو التجارية"
        rows_html = ""
        for idx, r in employees_df.iterrows():
            rows_html += f"""
            <tr>
                <td style="text-align:center; padding:7px; border:1px solid #CBD5E1;">{idx+1}</td>
                <td style="text-align:center; padding:7px; border:1px solid #CBD5E1;">{r['كود الموظف']}</td>
                <td style="text-align:right; padding:7px; border:1px solid #CBD5E1;">{r['اسم الموظف']}</td>
                <td style="text-align:center; padding:7px; border:1px solid #CBD5E1;">{r['المسمى الوظيفي']}</td>
                <td style="text-align:center; padding:7px; border:1px solid #CBD5E1;">{r['الراتب الإجمالي']:,.2f}</td>
                <td style="text-align:center; padding:7px; border:1px solid #CBD5E1;">{r['حصة الموظف تأمينات']:,.2f}</td>
                <td style="text-align:center; padding:7px; border:1px solid #CBD5E1; font-weight:bold;">{r['صافي الراتب الشهري']:,.2f}</td>
                <td style="text-align:center; padding:7px; border:1px solid #CBD5E1;">{r['قسط نهاية الخدمة الشهري']:,.2f}</td>
            </tr>
            """
        return f"""
        <!DOCTYPE html>
        <html lang="ar" dir="rtl">
        <head>
            <meta charset="UTF-8">
            <title>مسير الرواتب - {payroll.payroll_number}</title>
            <style>
                * {{ box-sizing: border-box; font-family: 'Cairo', 'Segoe UI', Tahoma, sans-serif; }}
                body {{ margin: 0; padding: 18px; background: #F8FAFC; color: #0F172A; }}
                .p-box {{ max-width: 860px; margin: 0 auto; background: white; padding: 26px; border-radius: 10px; border: 1px solid #CBD5E1; }}
                table {{ width: 100%; border-collapse: collapse; margin-top: 15px; font-size: 12.5px; }}
                th {{ background: #1E3A8A; color: white; padding: 8px; border: 1px solid #1E3A8A; }}
                .print-btn {{ background: #1E3A8A; color: white; border: none; padding: 10px; width: 100%; border-radius: 6px; font-weight: bold; cursor: pointer; margin-bottom: 14px; }}
                @media print {{ .no-print {{ display: none !important; }} body {{ background: white; padding: 0; }} }}
            </style>
        </head>
        <body>
            <div class="p-box">
                <button class="print-btn no-print" onclick="window.print()">🖨️ طباعة كشف مسير الرواتب المعتمد</button>
                <h2 style="margin:0; color:#1E3A8A;">{company_name} - كشف مسير الرواتب والأجور الشهري</h2>
                <p style="margin:5px 0; font-size:13px;"><b>رقم المسير:</b> {payroll.payroll_number} | <b>عن شهر:</b> {payroll.period_month} | <b>تاريخ الصرف:</b> {payroll.payment_date}</p>
                <table>
                    <thead>
                        <tr><th>#</th><th>الكود</th><th>اسم الموظف</th><th>المسمى</th><th>إجمالي الراتب</th><th>استقطاع التأمينات</th><th>صافي الراتب</th><th>مخصص نهاية الخدمة</th></tr>
                    </thead>
                    <tbody>{rows_html}</tbody>
                </table>
                <div style="margin-top:16px; background:#F1F5F9; padding:12px; border-radius:8px; display:flex; justify-content:space-around; font-size:13.5px;">
                    <span><b>إجمالي الرواتب:</b> {payroll.total_gross_salaries:,.2f} ريال</span>
                    <span><b>حصة الشركة تأمينات:</b> {payroll.total_gosi_employer:,.2f} ريال</span>
                    <span><b>مخصص نهاية الخدمة:</b> {payroll.total_eos_accrual:,.2f} ريال</span>
                    <span style="color:#15803D;"><b>صافي المنصرف للموظفين:</b> {payroll.net_paid_salaries:,.2f} ريال</span>
                </div>
            </div>
        </body>
        </html>
        """

    @staticmethod
    def generate_journal_voucher_html(
        entry: JournalEntry,
        settings: Optional[CompanySettings] = None
    ) -> str:
        """توليد سند قيد يومية رسمي قابل للطباعة والتوقيع (إعداد / مراجعة / اعتماد)"""
        company_name = settings.company_name_ar if settings else "شركة ميزان برو التجارية"
        company_vat = settings.vat_number if settings else "300000000000003"
        status_label = "معتمد ومرحل بالدفاتر ✅" if entry.status.value == "POSTED" else ("مسودة بانتظار الاعتماد ⏳" if entry.status.value == "DRAFT" else "قيد معكوس 🔄")

        rows_html = ""
        for jl in entry.lines:
            acc_code = jl.account.code if jl.account else "-"
            acc_name = jl.account.name_ar if jl.account else "-"
            cc_name = jl.cost_center.name_ar if jl.cost_center else "-"
            rows_html += f"""
            <tr>
                <td style="text-align:center; padding:8px; border:1px solid #CBD5E1;">{jl.line_number}</td>
                <td style="text-align:center; padding:8px; border:1px solid #CBD5E1; font-family:monospace; font-weight:bold;">{acc_code}</td>
                <td style="text-align:right; padding:8px; border:1px solid #CBD5E1;">{acc_name}</td>
                <td style="text-align:left; padding:8px; border:1px solid #CBD5E1; font-weight:bold; color:#1E3A8A;">{jl.debit:,.2f}</td>
                <td style="text-align:left; padding:8px; border:1px solid #CBD5E1; font-weight:bold; color:#15803D;">{jl.credit:,.2f}</td>
                <td style="text-align:center; padding:8px; border:1px solid #CBD5E1;">{cc_name}</td>
                <td style="text-align:right; padding:8px; border:1px solid #CBD5E1;">{html.escape(str(jl.description or entry.description or ""))}</td>
            </tr>
            """

        return f"""
        <!DOCTYPE html>
        <html lang="ar" dir="rtl">
        <head>
            <meta charset="UTF-8">
            <title>سند قيد يومية - {entry.entry_number}</title>
            <style>
                * {{ box-sizing: border-box; font-family: 'Cairo', 'Segoe UI', Tahoma, sans-serif; }}
                body {{ margin: 0; padding: 18px; background: #F8FAFC; color: #0F172A; }}
                .jv-box {{ max-width: 860px; margin: 0 auto; background: white; padding: 28px; border-radius: 12px; border: 1px solid #CBD5E1; }}
                .jv-header {{ display: flex; justify-content: space-between; align-items: center; border-bottom: 2px solid #1E3A8A; padding-bottom: 14px; margin-bottom: 16px; }}
                .badge {{ background: #1E3A8A; color: white; padding: 7px 18px; border-radius: 6px; font-weight: bold; font-size: 15px; }}
                .meta-grid {{ display: grid; grid-template-columns: 1fr 1fr; gap: 12px; background: #F1F5F9; padding: 12px 16px; border-radius: 8px; margin-bottom: 16px; font-size: 13px; }}
                table {{ width: 100%; border-collapse: collapse; font-size: 12.5px; margin-bottom: 16px; }}
                th {{ background: #1E3A8A; color: white; padding: 9px; border: 1px solid #1E3A8A; text-align: center; }}
                .print-btn {{ background: #1E3A8A; color: white; border: none; padding: 10px; width: 100%; border-radius: 6px; font-weight: bold; cursor: pointer; margin-bottom: 14px; }}
                @media print {{ .no-print {{ display: none !important; }} body {{ background: white; padding: 0; }} .jv-box {{ border: none; }} }}
            </style>
        </head>
        <body>
            <div class="jv-box">
                <button class="print-btn no-print" onclick="window.print()">🖨️ طباعة سند القيد اليومي الرسمي (أو الحفظ كـ PDF)</button>
                <div class="jv-header">
                    <div>
                        <h2 style="margin:0 0 4px 0; color:#1E3A8A;">{company_name}</h2>
                        <p style="margin:0; font-size:12px; color:#475569;">الرقم الضريبي: {company_vat}</p>
                    </div>
                    <div style="text-align:left;">
                        <div class="badge">سند قيد يومية عام (Journal Voucher)</div>
                        <p style="margin:4px 0 0 0; font-size:12px; color:#475569;">{status_label}</p>
                    </div>
                </div>
                <div class="meta-grid">
                    <div>
                        <p style="margin:3px 0;"><b>رقم القيد:</b> <span style="font-family:monospace; font-weight:bold;">{entry.entry_number}</span></p>
                        <p style="margin:3px 0;"><b>تاريخ القيد:</b> {entry.entry_date}</p>
                    </div>
                    <div>
                        <p style="margin:3px 0;"><b>نوع المستند:</b> {entry.doc_type.value}</p>
                        <p style="margin:3px 0;"><b>المرجع:</b> {entry.reference or "-"}</p>
                    </div>
                </div>
                <p style="font-size:13.5px; margin:0 0 12px 0;"><b>الشرح والبيان العام:</b> {html.escape(str(entry.description or ""))}</p>
                <table>
                    <thead>
                        <tr><th>#</th><th>كود الحساب</th><th>اسم الحساب</th><th>مدين (ريال)</th><th>دائن (ريال)</th><th>مركز التكلفة</th><th>البيان التفصيلي</th></tr>
                    </thead>
                    <tbody>
                        {rows_html}
                        <tr style="background:#FEF08A; font-weight:bold;">
                            <td colspan="3" style="text-align:center; padding:8px; border:1px solid #CBD5E1;">الإجمالي المتزن</td>
                            <td style="text-align:left; padding:8px; border:1px solid #CBD5E1;">{entry.total_debit:,.2f}</td>
                            <td style="text-align:left; padding:8px; border:1px solid #CBD5E1;">{entry.total_credit:,.2f}</td>
                            <td colspan="2" style="text-align:center; padding:8px; border:1px solid #CBD5E1;">متزن 100% ✅</td>
                        </tr>
                    </tbody>
                </table>
                <div style="margin-top:28px; display:flex; justify-content:space-between; font-size:13px; border-top:1px dashed #CBD5E1; padding-top:16px;">
                    <div><b>إعداد المحاسب:</b> ........................</div>
                    <div><b>مراجعة وتدقيق:</b> ........................</div>
                    <div><b>اعتماد المدير المالي:</b> ........................</div>
                </div>
            </div>
        </body>
        </html>
        """

    @staticmethod
    def generate_immutable_audit_pdf(db: Session, **kwargs) -> bytes:
        """توليد ملف PDF مؤمّن وغير قابل للتعديل لسجل التدقيق الرقابي للمراجع الخارجي"""
        from zatca_audit_pdf import ZATCAAuditPDFExporter
        return ZATCAAuditPDFExporter.generate_immutable_audit_pdf(db=db, **kwargs)

    @staticmethod
    def generate_audit_log_printable_html(db: Session, **kwargs) -> str:
        """توليد قالب المعاينة والطباعة الرسمي لسجل التدقيق الرقابي للمراجع الخارجي"""
        from zatca_audit_pdf import ZATCAAuditPDFExporter
        return ZATCAAuditPDFExporter.generate_printable_audit_html(db=db, **kwargs)

